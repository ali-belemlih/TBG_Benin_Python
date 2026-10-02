import os
import psycopg2
import re
import numpy as np
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook
from dotenv import load_dotenv
from sqlalchemy import create_engine
from urllib.parse import quote_plus
from dateutil.relativedelta import relativedelta

load_dotenv()

# Database connection details (modify as needed)
DB_CONFIG = {
    "dbname": os.getenv("DB_NAME"),
    "user": os.getenv("DB_USER"),
    "password": os.getenv("DB_PASSWORD"),
    "host": os.getenv("DB_HOST"),
    "port": os.getenv("DB_PORT"),
}


# Database Engine
def get_db_engine():
    user = os.getenv("DB_USER")
    password = quote_plus(os.getenv("DB_PASSWORD"))  # Encode special chars
    host = os.getenv("DB_HOST")
    port = os.getenv("DB_PORT")
    dbname = os.getenv("DB_NAME")

    uri = f"postgresql+psycopg2://{user}:{password}@{host}:{port}/{dbname}"
    return create_engine(uri)


# Database connection function
def get_db_connection():
    return psycopg2.connect(**DB_CONFIG)


# Generic ID lookup


def get_id_from_table(cur, table_name, name, foreign_key=None, foreign_id=None):
    if not name:
        return None

    query = f"SELECT id FROM {table_name} WHERE name ilike %s"
    params = [name]

    if foreign_key and foreign_id:
        query += f" AND {foreign_key} = %s"
        params.append(foreign_id)

    cur.execute(query + ";", tuple(params))
    result = cur.fetchone()
    return result[0] if result else None


# Category ID lookup with mapping


def get_financial_category_id(cur, name):
    category_mapping = {
        "P&L consolidé": "P&L conso",
        "Flux Financiers consolidés": "Cash Conso",
    }
    return get_id_from_table(
        cur, "financial_categories", category_mapping.get(name, name)
    )


# Hierarchical financial lookup


def get_financial_hierarchy(cur, name, category_id):
    tables = [
        ("financial_types", "financial_category_id"),
        ("financial_metric", "financial_type_id"),
        ("financial_submetric", "financial_metric_id"),
    ]
    foreign_id = category_id
    ids = []
    for table, fk in tables:
        current_id = get_id_from_table(cur, table, name, fk, foreign_id)
        ids.append(current_id)
        if current_id:
            foreign_id = current_id
        else:
            break

    return tuple((ids + [None, None, None])[:3])  # Ensure 3-tuple


# Generic conditional upsert for any table


def conditional_upsert(cur, table_name, key_fields: dict, update_fields: dict):
    where_clauses = []
    where_values = []

    for col, val in key_fields.items():
        if val is None:
            where_clauses.append(f"{col} IS NULL")
        else:
            where_clauses.append(f"{col} = %s")
            where_values.append(val)

    where_str = " AND ".join(where_clauses)

    cur.execute(f"SELECT 1 FROM {table_name} WHERE {where_str}", tuple(where_values))
    exists = cur.fetchone()
    if exists:
        if update_fields:
            update_clause = [f"{k} = %s" for k in update_fields]
            update_values = list(update_fields.values())
            cur.execute(
                f"""
                UPDATE {table_name}
                SET {", ".join(update_clause)}
                WHERE {where_str}
            """,
                tuple(update_values + where_values),
            )
            print(
                f"🔄 Updated in '{table_name}': Keys={key_fields}, Updates={update_fields}"
            )
    else:
        columns = list(key_fields.keys()) + list(update_fields.keys())
        values = list(key_fields.values()) + list(update_fields.values())
        placeholders = ["%s"] * len(values)
        cur.execute(
            f"""
            INSERT INTO {table_name} ({", ".join(columns)})
            VALUES ({", ".join(placeholders)})
        """,
            tuple(values),
        )
        print(
            f"➕ Inserted into '{table_name}': Keys={key_fields}, Values={update_fields}"
        )


# Specific upserts using shared logic
def upsert_financial_data(
    cur,
    table_name,
    type_id,
    metric_id,
    submetric_id,
    date_value,
    version_id,
    budget_value=None,
    real_value=None,
    actual1_value=None,
    actual2_value=None,
    actual3_value=None,
    last_year_real_value=None,
    adjusted_value=None
):
    key_fields = {
        "date": date_value,
        "financial_type_id": type_id,
        "financial_metric_id": metric_id,
        "financial_submetric_id": submetric_id,
        "version_id": version_id,
    }
    update_fields = {}
    if budget_value is not None:
        update_fields["budget_value"] = budget_value
    if real_value is not None:
        update_fields["real_value"] = real_value
    if actual1_value is not None:
        update_fields["actual1_value"] = actual1_value
    if actual2_value is not None:
        update_fields["actual2_value"] = actual2_value
    if actual3_value is not None:
        update_fields["actual3_value"] = actual3_value
    if last_year_real_value is not None:
        update_fields["last_year_real_value"] = last_year_real_value
    if adjusted_value is not None:
        update_fields["adjusted_value"] = adjusted_value
    conditional_upsert(cur, table_name, key_fields, update_fields)

def upsert_annual_data(
    cur,
    table_name,
    type_id,
    metric_id,
    submetric_id,
    date_value,
    version_id,
    budget_value=None,
    real_value=None,
    actual1_value=None,
    actual2_value=None,
    actual3_value=None,
    last_year_real_value=None,
):
    key_fields = {
        "date": date_value,
        "financial_type_id": type_id,
        "financial_metric_id": metric_id,
        "financial_submetric_id": submetric_id,
        "version_id": version_id,
    }
    update_fields = {}
    if budget_value is not None:
        update_fields["budget_value"] = budget_value
    if real_value is not None:
        update_fields["real_value"] = real_value
    if actual1_value is not None:
        update_fields["actual1_value"] = actual1_value
    if actual2_value is not None:
        update_fields["actual2_value"] = actual2_value
    if actual3_value is not None:
        update_fields["actual3_value"] = actual3_value
    if last_year_real_value is not None:
        update_fields["last_year_real_value"] = last_year_real_value
    conditional_upsert(cur, table_name, key_fields, update_fields)


def upsert_collapse_financial_data(
    cur,
    table_name,
    entity_id=None,
    entity_type=None,
    date_value=None,
    version_id=None,
    budget_value=None,
    real_value=None,
    actual1_value=None,
    actual2_value=None,
    actual3_value=None,
    last_year_real_value=None,
    adjusted_value = None
):
    if not all([entity_id, entity_type, date_value, version_id]):
        return None

    key_fields = {
        "entity_id": entity_id,
        "entity_type": entity_type,
        "date": date_value,
        "version_id": version_id,
    }
    update_fields = {}
    if budget_value is not None:
        update_fields["budget_value"] = budget_value
    if real_value is not None:
        update_fields["real_value"] = real_value
    if actual1_value is not None:
        update_fields["actual1_value"] = actual1_value
    if actual2_value is not None:
        update_fields["actual2_value"] = actual2_value
    if actual3_value is not None:
        update_fields["actual3_value"] = actual3_value
    if last_year_real_value is not None:
        update_fields["last_year_real_value"] = last_year_real_value
    if adjusted_value is not None:
        update_fields["adjusted_value"] = adjusted_value
    conditional_upsert(cur, table_name, key_fields, update_fields)


def upsert_cashflow_month_data(
    cur, entity_id, entity_type, year, month, value, version_id, current_year_total=None
):
    valid_months = [
        "jan", "feb", "mar", "apr", "may", "jun",
        "jul", "aug", "sep", "oct", "nov", "dec",
    ]
    valid_adjusted_months = [f"adjusted_{m}" for m in valid_months]

    month = month.lower()
    if month not in valid_months and month not in valid_adjusted_months:
        raise ValueError(
            f"Invalid month: {month}. Must be one of {valid_months} or {valid_adjusted_months}"
        )

    key_fields = {
        "entity_id": entity_id,
        "entity_type": entity_type,
        "year": year,
        "version_id": version_id,
    }
    update_fields = {month: value}
    if current_year_total is not None:
        update_fields["current_year_total"] = current_year_total

    conditional_upsert(cur, "cashflow_data", key_fields, update_fields)

def upsert_collapse_monthly_data(
    cur, table_name, entity_id, entity_type, date_value, version_id, real_value=None
):
    if not all([entity_id, entity_type, date_value, version_id]):
        return None

    key_fields = {
        "entity_id": entity_id,
        "entity_type": entity_type,
        "date": date_value,
        "version_id": version_id,
    }

    # Fetch current value
    cur.execute(
        f"""
        SELECT real_value FROM {table_name}
        WHERE entity_id = %s AND entity_type = %s AND date = %s AND version_id = %s
    """,
        (entity_id, entity_type, date_value, version_id),
    )
    result = cur.fetchone()

    if result:
        current_value = result[0] or 0.0
        new_value = current_value + (real_value or 0.0)
    else:
        new_value = real_value

    update_fields = {"real_value": new_value}
    conditional_upsert(cur, table_name, key_fields, update_fields)


def upsert_cumulative_data(conn, cumulative_data, version_id):
    """Insert or update records into financial_cumulative_data table."""

    print(f"Upserting cumulative data for {len(cumulative_data)} records")

    with conn.cursor() as cur:
        for row in cumulative_data:
            (
                financial_type_id,
                financial_metric_id,
                financial_submetric_id,
                date,
                real_value,
                budget_value,
                last_year_real_value,
                actual1_value,
                actual2_value,
                actual3_value,
            ) = row

            key_fields = {
                "financial_type_id": financial_type_id,
                "financial_metric_id": financial_metric_id,
                "financial_submetric_id": financial_submetric_id,
                "date": date,
                "version_id": version_id,
            }

            update_fields = {
                "real_value": real_value,
                "budget_value": budget_value,
                "last_year_real_value": last_year_real_value,
                "actual1_value": actual1_value,
                "actual2_value": actual2_value,
                "actual3_value": actual3_value,
            }

            conditional_upsert(
                cur, "public.financial_cumulative_data", key_fields, update_fields
            )

        conn.commit()
    print("Cumulative data inserted/updated successfully")


def upsert_collapse_cumulative_data(conn, cumulative_data, version_id):
    """Upsert records into collapse_cumul_data table with version_id."""

    print(f"Upserting cumulative data for {len(cumulative_data)} records")

    with conn.cursor() as cur:
        for row in cumulative_data:
            key_fields = {
                "entity_id": row[0],
                "entity_type": row[1],
                "date": row[2],
                "version_id": version_id,
            }
            update_fields = {
                "real_value": row[3],
                "budget_value": row[4],
                "last_year_real_value": row[5],
                "actual1_value": row[6],
                "actual2_value": row[7],
                "actual3_value": row[8],
            }
            conditional_upsert(cur, "collapse_cumul_data", key_fields, update_fields)
        conn.commit()
    print("Collapsed cumulative data inserted/updated successfully")


# Fetch version id via version_name
def get_version_id_by_name(version_name):
    try:
        conn = get_db_connection()
        with conn:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT id FROM tbg_version WHERE version_name = %s",
                    (version_name,),
                )
                result = cur.fetchone()

                if result:
                    return result[0]
                else:
                    raise ValueError(f"No version found for name: {version_name}")
    except Exception as e:
        raise RuntimeError(f"Database error while fetching version_id: {e}")


def get_month_year_by_version_name(version_name):
    try:
        conn = get_db_connection()
        with conn:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT month, year FROM tbg_version WHERE version_name = %s",
                    (version_name,),
                )
                result = cur.fetchone()

                if result:
                    month, year = result
                    return f"{int(month):02d}{int(year)}"
                else:
                    raise ValueError(f"No version found for name: {version_name}")
    except Exception as e:
        raise RuntimeError(f"Database error while fetching month and year: {e}")


# Value fetch by type


def get_real_or_budget_value(cur, source_info, date, label, version_id):
    table = source_info.get("table_name")
    value = source_info.get("value")
    entity_type = source_info.get("entity_type")
    entity_id = source_info.get("entity_id")
    if isinstance(date, str):
        date = datetime.strptime(date, "%Y-%m-%d")
    if "month" in source_info:
        date = date - relativedelta(months=1)
        if date.month == 12:
            version_id = get_default_version_id(date.month, date.year)

    print(f"================Month: {date}")
    if not table:
        return 0.0

    if isinstance(version_id, np.generic):
        version_id = version_id.item()

    if "column_name" in source_info:
        return _get_value_by_column_lookup(
            cur, table, source_info["column_name"], value, date, label, version_id
        )
    elif "entity_type" in source_info:
        return _get_value_by_entity_month(
            cur, table, entity_type, entity_id, date, version_id, label
        )

    print(f"Invalid source_info: {source_info}")
    return 0.0


def _get_value_by_column_lookup(cur, table, column, value, date, label, version_id):
    if not column:
        return 0.0

    # Only add adjusted_value when fetching real_value
    if label == "real_value":
        query = f"""
            SELECT COALESCE({label}, 0.0) + COALESCE(adjusted_value, 0.0)
            FROM {table}
            WHERE {column} = %s AND date = %s AND version_id = %s
            LIMIT 1
        """
    else:
        query = f"""
            SELECT COALESCE({label}, 0.0)
            FROM {table}
            WHERE {column} = %s AND date = %s AND version_id = %s
            LIMIT 1
        """

    cur.execute(query, (value, date, version_id))
    result = cur.fetchone()
    return result[0] if result else 0.0


def _get_value_by_entity_month(cur, table, entity_type, value, date, version_id, label):
    if not entity_type:
        return 0.0

    if table == "collapse_monthly_data" or table == 'collapse_cumul_data':
        # Only add adjusted_value when fetching real_value
        if label == "real_value":
            query = f"""
                SELECT COALESCE({label}, 0.0) + COALESCE(adjusted_value, 0.0)
                FROM {table}
                WHERE entity_type = %s AND entity_id = %s
                AND date = %s AND version_id = %s
                LIMIT 1
            """
        else:
            query = f"""
                SELECT COALESCE({label}, 0.0)
                FROM {table}
                WHERE entity_type = %s AND entity_id = %s
                AND date = %s AND version_id = %s
                LIMIT 1
            """
        cur.execute(query, (entity_type, value, date, version_id))
        result = cur.fetchone()
        return result[0] if result else 0.0

    try:
        if isinstance(date, str):
            date = datetime.strptime(date, "%Y-%m-%d")
        month = date.month
    except Exception as e:
        raise ValueError(f"❌ Failed to parse date: {e}")

    column_name = {
        1: "jan", 2: "feb", 3: "mar", 4: "apr",
        5: "may", 6: "jun", 7: "jul", 8: "aug",
        9: "sep", 10: "oct", 11: "nov", 12: "dec",
    }.get(month, "jan")

    adjusted_column = f"adjusted_{column_name}"

    print(f"========column_name: {column_name} & table: {table} & entity_type: {entity_type} & entity_id: {value} & version_id: {version_id}")

    # cashflow_data month columns always get adjusted added
    query = f"""
        SELECT COALESCE({column_name}, 0.0) + COALESCE({adjusted_column}, 0.0)
        FROM {table}
        WHERE entity_type = %s AND entity_id = %s
        AND version_id = %s
        LIMIT 1
    """
    cur.execute(query, (entity_type, value, version_id))
    result = cur.fetchone()
    return result[0] if result and result[0] is not None else 0.0

def get_default_version_id(month: str, year: int, engine=None):
    if engine is None:
        engine = get_db_engine()

    query = f"""
        SELECT id FROM tbg_version
        WHERE month = '{month}' AND year = '{year}'
        AND is_default = true;
    """
    result = pd.read_sql_query(query, engine)
    return result["id"].iloc[0] if not result.empty else None


def get_sage_real_value(cur, sage_version, source_key, category=None):
    txsnam = category or "YYOPEX"

    cur.execute(
        """
        SELECT
            SUM(CAST(real_value AS FLOAT)) as total_value
        FROM (
            SELECT
                MAX(CASE WHEN COL_0 = '0' THEN SPLIT_PART(AMTVAL_0, '-', 1) END) AS sage_source_key,
                MAX(CASE WHEN COL_0 = '1' THEN AMTVAL_0 END) AS real_value,
                LIG_0
            FROM sage_yexptdb
            WHERE VERSION_0 = 'YEXPTDB'
              AND IND_0 = '0'
              AND TXSNAM_0 = %s
              AND COL_0 IN ('0', '1')
              AND sage_version = %s
            GROUP BY LIG_0
        ) AS subquery
        WHERE sage_source_key = %s;
    """,
        (txsnam, sage_version, source_key),
    )

    row = cur.fetchone()
    return float(row[0]) if row and row[0] is not None else 0.0


def get_calculation_field_value(cur, tbg_key, date, field_type, version_id):
    parsed_date = datetime.strptime(date, "%Y-%m-%d").date()
    month = parsed_date.month
    year = parsed_date.year

    try:
        cur.execute(
            """
            SELECT cf.input_value
            FROM tbg_calculation_fields cf
            JOIN tbg_calculation_fields_types ccf
              ON cf.input_field_id = ccf.id
            WHERE ccf.tbg_key = %s
              AND ccf.key = %s
              AND cf.tbg_version_id = %s
              AND cf.month = %s
              AND cf.year = %s
        """,
            (tbg_key, field_type, version_id, month, year),
        )

        row = cur.fetchone()
        if row:
            return float(row[0])
        else:
            print(
                f"⚠️ No calculation field found for tbg_key={tbg_key}, type={field_type}, version_id={version_id}, date={date}"
            )
            return 0.0

    except Exception as e:
        print(
            f"❌ Failed to fetch calculation field value for {tbg_key}, type={field_type}, version_id={version_id}: {e}"
        )
        return 0.0
