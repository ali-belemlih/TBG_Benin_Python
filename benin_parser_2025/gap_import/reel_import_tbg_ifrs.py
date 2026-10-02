from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from helpers.parse_arg import parse_arguments
from helpers.db_utils import get_db_connection, get_version_id_by_name
from services.minio_factory import get_minio_service

# Setup
conn = get_db_connection()
cur = conn.cursor()

month_mapping = {
    "01": "JAN", "02": "FEV", "03": "MAR", "04": "AVRIL",
    "05": "MAI", "06": "JUIN", "07": "JUIL", "08": "AOÛT",
    "09": "SEPTEMBRE", "10": "OCTOBRE", "11": "NOVEMBRE", "12": "DÉCEMBRE"
}

col_mapping = {
    "JAN": "AN", "FEV": "AO", "MAR": "AP", "AVRIL": "AQ", "MAI": "AR",
    "JUIN": "AS", "JUIL": "AT", "AOÛT": "AU", "SEPTEMBRE": "AV",
    "OCTOBRE": "AW", "NOVEMBRE": "AX", "DÉCEMBRE": "AY"
}


def fetch_with_headers(query, params):
    cur.execute(query, params)
    columns = [desc[0] for desc in cur.description]
    rows = cur.fetchall()
    return [dict(zip(columns, row)) for row in rows]


def fetch_financial_submetric_data(tbg_key, date):
    query = """
    SELECT fmd.date, fmd.real_value, fsm.id AS financial_submetric_id,
           fsm.name AS financial_submetric_name, fsm.tbg_key, fsm.sage_source_key
    FROM financial_metrics_data fmd
    LEFT JOIN financial_submetric fsm ON fmd.financial_submetric_id = fsm.id
    WHERE fsm.tbg_key = %s AND fmd.date = %s
    ORDER BY fmd.date, fsm.name;
    """
    return fetch_with_headers(query, (tbg_key, date))


def update_financial_submetric_real_value(tbg_key, date, real_value, version_id):
    update_query = """
    UPDATE financial_metrics_data fmd
    SET real_value = %s, version_id = %s
    FROM financial_submetric fsm
    WHERE fmd.financial_submetric_id = fsm.id
      AND fsm.tbg_key = %s AND fmd.date = %s
    """
    cur.execute(update_query, (real_value, version_id, tbg_key, date))
    conn.commit()
    return cur.rowcount


def get_opposite_value_from_sheet(wb, month_label, row=52, col='B', sheet_names=("P&amp;L conso", "P&L conso")):
    for sheet_name in sheet_names:
        if sheet_name in wb:
            ws = wb[sheet_name]
            month_col = col_mapping.get(month_label)
            if month_col:
                value = ws[f"{month_col}{row}"].value
                return -1 * value
    return None


def process_month_data(wb, tbg_key, year, month, version_id):
    date = f"{year}-{month}-01"
    month_label = month_mapping[month]
    opposite_value = get_opposite_value_from_sheet(wb, month_label)

    if opposite_value is not None:
        print(f"Processing {date} | Opposite value: {opposite_value}")
        before_update = fetch_financial_submetric_data(tbg_key, date)
        print(f"Before update: {before_update[0]['real_value'] if before_update else 'Not found'}")

        update_financial_submetric_real_value(tbg_key, date, opposite_value, version_id)

        after_update = fetch_financial_submetric_data(tbg_key, date)
        print(f"After update: {after_update[0]['real_value'] if after_update else 'Not found'}")
    else:
        print(f"Data not found in Excel for {month_label}.")


if __name__ == "__main__":
    args = parse_arguments()
    file_name = args.file_name
    tbg_key = 'PL46'

    target_month = args.month_year[:2]
    target_year = args.month_year[2:]
    version_id = get_version_id_by_name(args.version_id)
    object_path = f"{target_year}/{target_year}{target_month}/{file_name}"
    print(f"Fetching file from MinIO: {object_path}")

    minio_service = get_minio_service()
    file_bytes = minio_service.get_file_bytes(object_path)
    workbook = load_workbook(filename=BytesIO(file_bytes), data_only=True, read_only=True)

    process_month_data(workbook, tbg_key, target_year, target_month, version_id)
