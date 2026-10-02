import os
import json
import pandas as pd
from openpyxl import load_workbook
from collections import defaultdict
from helpers.parse_arg import parse_arguments
from helpers.export_utils import read_json_mapping, normalize_month_key, get_version_for_tbg_export, excel_column_name
from helpers.db_utils import get_db_engine


def extract_type_id_rows(raw_config):
    type_id_to_rows = defaultdict(list)
    for section in raw_config.values():
        if "collapsible_types" in section:
            for type_id_str, type_map in section["collapsible_types"].items():
                try:
                    type_id = int(type_id_str)
                    for _, row_number in type_map.items():
                        type_id_to_rows[type_id].append(row_number)
                except Exception as e:
                    print(f"⚠️ Error processing {type_id_str}: {e}")
    return type_id_to_rows


def run_sql_query(path, replacements, engine):
    with open(path, "r") as f:
        sql = f.read()
        for key, val in replacements.items():
            sql = sql.replace(key, val)
    return pd.read_sql_query(sql, engine)


def write_data_to_worksheet(df, worksheet, id_to_rows_map, start_col, sequence):
    for _, row in df.iterrows():
        id_value = row.get("collapse_type_id", 0) or 0
        row_numbers = id_to_rows_map.get(id_value, [])

        for row_number in row_numbers:
            for i, column_name in enumerate(sequence):
                start_col_index = sum((ord(c) - 64) * (26 ** j) for j, c in enumerate(reversed(start_col)))
                col_letter = excel_column_name(start_col_index + i)
                cell_address = f"{col_letter}{row_number}"

                if column_name in row and pd.notnull(row[column_name]):
                    worksheet[cell_address] = row[column_name]


def process_sheet(wb, engine, base_path, sheet_name, version_id, month_param, year_param):
    print(f"🔄 Processing sheet: {sheet_name}")
    ws = wb[sheet_name]

    mapping_base = f"{base_path}/mapping/ca_mobile"
    sql_base = f"{base_path}/sql_files"

    collapsible_mapping = read_json_mapping(f"{mapping_base}/ca_mobile_collapsible_types.json")
    month_mapping = read_json_mapping(f"{mapping_base}/month_mapping.json")
    category_mapping = read_json_mapping(f"{mapping_base}/ca_mobile_collapsible_categories.json")
    subcategory_mapping = read_json_mapping(f"{mapping_base}/ca_mobile_collapsible_subcategories.json")
    annual_mapping_path = f"{mapping_base}/annual_mapping.json"
    cumul_mapping_path = f"{mapping_base}/ca_cumul_mapping.json"

    type_id_to_rows = extract_type_id_rows(collapsible_mapping)

    for month_key, details in month_mapping.items():
        normalized_month = normalize_month_key(month_key)
        version = get_version_for_tbg_export(normalized_month, month_param, year_param, version_id, engine)

        query_date = f"{year_param}-{details['date'][5:7]}-{details['date'][8:]}"
        loop_month = details['date'][5:7]
        if int(loop_month) > int(month_param):
            continue
        replacements = {"{{date}}": query_date, "{version_id}": str(version)}

        # Process items
        df = run_sql_query(f"{sql_base}/ca_collapsible_items_query.sql", replacements, engine)
        write_data_to_worksheet(df, ws, type_id_to_rows, details["start_col"], details["sequence"])

        # Process categories
        df_cat = run_sql_query(f"{sql_base}/ca_collapsible_categories_query.sql", replacements, engine)
        for _, row in df_cat.iterrows():
            row_number = category_mapping.get(str(row["category_id"]), {}).get(row["category_name"])
            if row_number:
                write_data_to_worksheet(pd.DataFrame([row]), ws, {0: [row_number]}, details["start_col"], details["sequence"])

        # Process subcategories
        if os.path.exists(f"{sql_base}/ca_collapsible_subcategories_query.sql"):
            df_sub = run_sql_query(f"{sql_base}/ca_collapsible_subcategories_query.sql", replacements, engine)
            for _, row in df_sub.iterrows():
                row_number = subcategory_mapping.get(str(row["subcategory_id"]), {}).get(row["subcategory_name"])
                if row_number:
                    write_data_to_worksheet(pd.DataFrame([row]), ws, {0: [row_number]}, details["start_col"], details["sequence"])

        # Annual data
        if os.path.exists(annual_mapping_path):
            annual_mapping = read_json_mapping(annual_mapping_path)
            if "annual" in annual_mapping:
                details = annual_mapping["annual"]
                query_date = f"{year_param}-01-01"
                replacements = {"2024-01-01": query_date, "{version_id}": str(version)}

                # Items
                df_annual = run_sql_query(f"{sql_base}/collapsible_annual_query.sql", replacements, engine)
                write_data_to_worksheet(df_annual, ws, type_id_to_rows, details["start_col"], details["sequence"])

                # Categories
                if os.path.exists(f"{sql_base}/ca_collapse_annual_data_categories_query.sql"):
                    df_annual_cat = run_sql_query(f"{sql_base}/ca_collapse_annual_data_categories_query.sql", replacements, engine)
                    for _, row in df_annual_cat.iterrows():
                        row_number = category_mapping.get(str(row["category_id"]), {}).get(row["category_name"])
                        if row_number:
                            write_data_to_worksheet(pd.DataFrame([row]), ws, {0: [row_number]}, details["start_col"], details["sequence"])

                # Subcategories
                if os.path.exists(f"{sql_base}/ca_collapsible_annual_subcategory_query.sql"):
                    df_annual_sub = run_sql_query(f"{sql_base}/ca_collapsible_annual_subcategory_query.sql", replacements, engine)
                    for _, row in df_annual_sub.iterrows():
                        row_number = subcategory_mapping.get(str(row["subcategory_id"]), {}).get(row["subcategory_name"])
                        if row_number:
                            write_data_to_worksheet(pd.DataFrame([row]), ws, {0: [row_number]}, details["start_col"], details["sequence"])
        # CUMULATIVE DATA
    normalized_month = normalize_month_key(month_key)
    version = get_version_for_tbg_export(normalized_month, month_param, year_param, version_id, engine)
    if os.path.exists(cumul_mapping_path):
            cumul_mapping = read_json_mapping(cumul_mapping_path)
            if "cumul" in cumul_mapping:
                details = cumul_mapping["cumul"]
                query_date = f"{year_param}-{month_param}-01"
                replacements = {"{{date}}": query_date, "{version_id}": str(version_id)}

                # Items
                df_cumul = run_sql_query(f"{sql_base}/collapsible_cumul_query.sql", replacements, engine)
                value = ""
                if 4<=int(month_param)<=5:
                    value = '1'
                elif 6<=int(month_param)<=8:
                    value = '2'
                elif 9<=int(month_param)<=12:
                    value = '3'
                sequence = [ ele.replace('#', value) if '#' in ele else ele for ele in details["sequence"]]
                write_data_to_worksheet(df_cumul, ws, type_id_to_rows, details["start_col"], sequence)

                # Categories
                if os.path.exists(f"{sql_base}/ca_mobile_collapsible_cumul_category.sql"):
                    df_cumul_cat = run_sql_query(f"{sql_base}/ca_mobile_collapsible_cumul_category.sql", replacements, engine)
                    for _, row in df_cumul_cat.iterrows():
                        row_number = category_mapping.get(str(row["category_id"]), {}).get(row["category_name"])
                        if row_number:
                            write_data_to_worksheet(pd.DataFrame([row]), ws, {0: [row_number]}, details["start_col"], sequence)

                # Subcategories
                if os.path.exists(f"{sql_base}/ca_mobile_collapsible_cumul_subcategory.sql"):
                    df_cumul_sub = run_sql_query(f"{sql_base}/ca_mobile_collapsible_cumul_subcategory.sql", replacements, engine)
                    for _, row in df_cumul_sub.iterrows():
                        row_number = subcategory_mapping.get(str(row["subcategory_id"]), {}).get(row["subcategory_name"])
                        if row_number:
                            write_data_to_worksheet(pd.DataFrame([row]), ws, {0: [row_number]}, details["start_col"], sequence)


def main():
    args = parse_arguments()
    month_year = args.month_year
    version_id = args.version_id
    month_param, year_param = month_year[:2], month_year[2:]

    base_path = "benin_tbg_export"
    output_path = f"{base_path}/outputs/tbg_report_{month_year}.xlsx"

    wb = load_workbook(output_path)
    engine = get_db_engine()

    sheet_name = "CA Mobile"
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in Excel template.")

    process_sheet(wb, engine, base_path, sheet_name, version_id, month_param, year_param)

    wb.save(output_path)
    print("📘 Report saved successfully.")

if __name__ == "__main__":
    main()
