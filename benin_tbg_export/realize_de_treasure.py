import os
import pandas as pd
from openpyxl.reader.excel import load_workbook
from helpers.parse_arg import parse_arguments
from helpers.export_utils import read_json_mapping, normalize_month_key, get_version_for_tbg_export, excel_column_name
from helpers.db_utils import get_db_engine


def load_sql_query(file_path, replacements=None):
    with open(file_path, "r") as f:
        query = f.read()
        if replacements:
            for key, value in replacements.items():
                query = query.replace(key, value)
        return query


def populate_section_data(ws, df, mapping, month_col_map, section_label):
    for _, row in df.iterrows():
        entity_id = str(row["entity_id"])
        if entity_id not in mapping:
            continue

        for designation, row_number in mapping[entity_id].items():
            # Last year total (col D)
            last_year_val = row['last_year_total']
            if pd.notnull(last_year_val) and last_year_val != 0:
                ws[f"{month_col_map['last_year']['sheet_col']}{row_number}"] = last_year_val

            # Current year total (col E) — written directly from DB, no formula needed
            current_year_val = row['current_year_total']
            if pd.notnull(current_year_val) and current_year_val != 0:
                ws[f"{month_col_map['current_year']['sheet_col']}{row_number}"] = current_year_val

            # Monthly data
            for month, col_info in month_col_map.items():
                if month in ('last_year', 'current_year'):  # skip non-month keys
                    continue
                col_letter = col_info.get("sheet_col", "A")
                value = row.get(
                    month.lower() if section_label in ("category", "subcategory") else month,
                    None
                )
                if pd.notnull(value) and value != 0:
                    ws[f"{col_letter}{row_number}"] = value

def replace_year_in_text(text, year_param):
    last_year, prev_year = str(int(year_param) - 1), str(int(year_param) - 2)
    temp = "___TEMP___"
    s = str(text).replace(prev_year, temp).replace(last_year, str(year_param)).replace(temp, last_year)
    return int(s) if isinstance(text, int) and s.isdigit() else s


def update_realize_report_header(ws, year_param):
    """
    Directly sets header values based on year_param, no chain replacement.
    D3: TOTAL {year-1}
    E3: TOTAL {year}
    F3: REALISATIONS {year}
    G3: REALISATIONS {year}
    """
    current_year = str(int(year_param))
    last_year = str(int(year_param) - 1)

    ws["D3"] = f"TOTAL {last_year}"
    ws["E3"] = f"TOTAL {current_year}"
    ws["F3"] = f"REALISATIONS\n{current_year}"
    ws["G3"] = f"REALISATIONS {current_year}"

    print(f"✅ Updated headers: D3='TOTAL {last_year}', E3='TOTAL {current_year}', F3='REALISATIONS {current_year}', G3='REALISATIONS {current_year}'")

def process_sheet(sheet_data, wb, engine, year_param, month_param, base_path, version_id):
    sheet_name = sheet_data["sheet_name"]

    print(f"Processing sheet: {sheet_name}")

    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Worksheet {sheet_name} not found in the template")

    ws = wb[sheet_name]

    # Load mappings
    config_data = {key: read_json_mapping(sheet_data[key]) for key in [
        "sheet_mapping_path", "month_mapping_path",
        "total_mapping_path", "subtotal_mapping_path",
        "category_mapping_path", "subcategory_mapping_path"
    ]}
    month_col_map = config_data["month_mapping_path"]

    # Total section

    total_query = load_sql_query(f"{base_path}/sql_files/realize_total_query.sql",
                                  {"{section_type}": "Total", "{version_id}": str(version_id)})
    totals_df = pd.read_sql_query(total_query, engine)
    populate_section_data(ws, totals_df, config_data["total_mapping_path"], month_col_map, "total")

    # Subtotal section
    subtotal_query = load_sql_query(f"{base_path}/sql_files/realize_total_query.sql",
                                     {"{section_type}": "Subtotal", "{version_id}": str(version_id)})
    subtotals_df = pd.read_sql_query(subtotal_query, engine)
    populate_section_data(ws, subtotals_df, config_data["subtotal_mapping_path"], month_col_map, "subtotal")

    # Category section
    category_query = load_sql_query(f"{base_path}/sql_files/realize_category_query.sql",
                                     {"{version_id}": str(version_id)})
    categories_df = pd.read_sql_query(category_query, engine)
    populate_section_data(ws, categories_df, config_data["category_mapping_path"], month_col_map, "category")

    # Subcategory section
    print("Fetching Subcategory Data...")
    subcategory_query = load_sql_query(f"{base_path}/sql_files/realize_subcategories_query.sql",
                                       {"{version_id}": str(version_id)})
    subcategories_df = pd.read_sql_query(subcategory_query, engine)
    populate_section_data(ws, subcategories_df, config_data["subcategory_mapping_path"], month_col_map, "subcategory")
    update_realize_report_header(ws, year_param)


def main():
    args = parse_arguments()
    month_year = args.month_year
    version_id = args.version_id
    month_param, year_param = month_year[:2], month_year[2:]

    base_path = "benin_tbg_export"
    output_path = f"{base_path}/outputs/tbg_report_{month_year}.xlsx"

    wb = load_workbook(output_path)
    sheets_to_process = read_json_mapping(f"{base_path}/configs/realize_de_treasure_config.json")
    engine = get_db_engine()

    try:
        print(f"🔄 Processing for {month_year} (version_id={version_id})")
        for sheet in sheets_to_process:
            process_sheet(sheet, wb, engine, year_param, month_param, base_path, version_id)

        wb.save(output_path)
        print("Financial report has been generated successfully!")

    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    main()
