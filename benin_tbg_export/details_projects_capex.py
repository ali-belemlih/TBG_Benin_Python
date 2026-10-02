import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Border, PatternFill, Alignment, Protection
from helpers.parse_arg import parse_arguments
from helpers.export_utils import excel_column_name
from helpers.db_utils import get_db_engine

def copy_row_style(ws, source_row, target_row):
    for col in range(1, ws.max_column + 1):
        source_cell = ws.cell(row=source_row, column=col)
        target_cell = ws.cell(row=target_row, column=col)

        if source_cell.has_style:
            target_cell.font = Font(**source_cell.font.__dict__)
            target_cell.border = Border(**source_cell.border.__dict__)
            target_cell.fill = PatternFill(**source_cell.fill.__dict__)
            target_cell.alignment = Alignment(**source_cell.alignment.__dict__)
            target_cell.number_format = source_cell.number_format
            target_cell.protection = Protection(**source_cell.protection.__dict__)


def write_project_data_to_sheet(ws, df_projects, df_financials, row_start=7):
    row_num = row_start
    rc = 0
    total_cumul_columns = ["I", "M", "Q", "U", "Y", "AC", "AG", "AK", "AO", "AS", "AW", "BA"]

    for _, project in df_projects.iterrows():
        row = [
            project["project_title"],
            project["contract_no"],
            project["contract_date"],
            project["supplier_name"],
            project["direction_name"]
        ]

        for month in range(1, 13):
            data = df_financials[
                (df_financials["capex_projects_id"] == project["id"]) &
                (df_financials["month"] == month)
            ].fillna(0)

            if not data.empty:
                equipment = data.iloc[0].get("equipment", 0)
                services = data.iloc[0].get("services", 0)
                add_costs = data.iloc[0].get("additional_costs", 0)
                total = equipment + services + add_costs
                row += [equipment, services, add_costs, total]
            else:
                row += [0, 0, 0, 0]

        # Add cumulative formula
        cumul_formula = "=" + "+".join([f"{col}{row_num}" for col in total_cumul_columns])
        row.append(cumul_formula)

        ws.insert_rows(row_num + 1)
        ws.row_dimensions[row_num + 1].hidden = False
        copy_row_style(ws, 7, row_num + 1)

        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_num, column=col_idx, value=value)

        # Update SUM formulas
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=9 + rc, column=col_idx)
            ws.row_dimensions[9 + rc].hidden = False
            if isinstance(cell.value, str) and "SUM" in cell.value:
                cell.value = f"=SUM({excel_column_name(col_idx)}7:{excel_column_name(col_idx)}{9 + rc - 2})"

        row_num += 1
        rc += 1

def replace_year_in_text(text, year_param):
    last_year, prev_year = str(int(year_param) - 1), str(int(year_param) - 2)
    s = str(text).replace(last_year, year_param)
    return int(s) if isinstance(text, int) and s.isdigit() else s

def update_details_project_header(ws, year_param):
    cols = ["F", "G", "N","R","V","Z","AD","AH","AL","AP","AT","AX"]
    for col in cols:
        val = ws[f"{col}4"].value
        if val:
            ws[f"{col}4"] = replace_year_in_text(val, year_param)
    cumul_val = ws["BB5"].value
    ws["BB5"] = replace_year_in_text(cumul_val, year_param)
    ws.merge_cells('BB4:BB5')
    ws['BB4'] = f'Capex mis en service {year_param}'
    ws['BB4'].alignment = Alignment(horizontal='center', vertical='center')
    ws['BB4'].font = Font(bold=True)

def process_capex_sheet(wb, engine, base_path, year):
    sheet_name = "Détail projets Capex"

    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Worksheet '{sheet_name}' not found in the template.")

    ws = wb[sheet_name]

    with open(f"{base_path}/sql_files/details_capex_projects_query.sql") as f:
        query1 = f.read().replace('{year_to_process}', year)
    with open(f"{base_path}/sql_files/capex_data_query.sql") as f:
        query2 = f.read().replace('{year_to_process}', year)

    df_projects = pd.read_sql(query1, engine)
    df_financials = pd.read_sql(query2, engine)

    write_project_data_to_sheet(ws, df_projects, df_financials)
    update_details_project_header(ws, year)


def main():
    args = parse_arguments()
    month_year = args.month_year
    version_id = args.version_id
    month, year = month_year[:2], month_year[2:]

    base_path = "benin_tbg_export"
    output_path = f"{base_path}/outputs/tbg_report_{month_year}.xlsx"
    engine = get_db_engine()

    try:
        wb = load_workbook(output_path)
        process_capex_sheet(wb, engine, base_path, year)
        wb.save(output_path)
        print("Financial report has been generated successfully!")
    except Exception as e:
        print(f"Error generating the report: {e}")


if __name__ == "__main__":
    main()
