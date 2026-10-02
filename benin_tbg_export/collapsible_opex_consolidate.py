import os
import json
import pandas as pd
from collections import defaultdict
from openpyxl import load_workbook
from helpers.parse_arg import parse_arguments
from helpers.export_utils import read_json_mapping, normalize_month_key, get_version_for_tbg_export, excel_column_name
from helpers.db_utils import get_db_engine


def load_sql_query(file_path, replacements=None):
    with open(file_path, "r") as f:
        query = f.read()
        if replacements:
            for key, val in replacements.items():
                query = query.replace(key, val)
        return query


def build_type_id_to_rows_map(raw_config):
    type_id_to_rows = defaultdict(list)

    def recurse_config(cfg):
        if isinstance(cfg, dict):
            for section in cfg.values():
                if "collapsible_types" in section:
                    for type_id_str, type_map in section["collapsible_types"].items():
                        try:
                            type_id = int(type_id_str)
                            type_id_to_rows[type_id].extend(type_map.values())
                        except Exception as e:
                            print(f"⚠️ Skipped invalid type ID '{type_id_str}': {e}")

    recurse_config(raw_config)
    return type_id_to_rows


def get_column_letter_index(col_str):
    return sum((ord(c) - 64) * (26 ** i) for i, c in enumerate(reversed(col_str)))


def populate_data_to_sheet(ws, df, type_id_to_rows, start_col, sequence):
    start_col_index = get_column_letter_index(start_col)

    for _, row in df.iterrows():
        type_id = row.get("collapse_type_id")
        row_numbers = type_id_to_rows.get(type_id, [])

        for row_num in row_numbers:
            for i, col_name in enumerate(sequence):
                col_letter = excel_column_name(start_col_index + i)
                value = row.get(col_name)

                if pd.notnull(value):
                    ws[f"{col_letter}{row_num}"] = value


def process_sheet(sheet_data, wb, engine, year_param, month_param, version_id, base_path):
    sheet_name = sheet_data["sheet_name"]

    print(f"Processing sheet: {sheet_name}")

    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Worksheet {sheet_name} not found in the template")

    ws = wb[sheet_name]

    # Load config and mappings
    raw_config = read_json_mapping(sheet_data["collapsible_type_mapping_path"])
    month_mapping = read_json_mapping(sheet_data["month_mapping_path"])
    annual_mapping_path = sheet_data.get("annual_mapping_path")
    cumul_mapping_path = sheet_data.get("cumul_mapping_path")

    type_id_to_rows = build_type_id_to_rows_map(raw_config)

    # --- Process Monthly Data ---
    for month_key, details in month_mapping.items():
        normalized_month = normalize_month_key(month_key)
        version = get_version_for_tbg_export(normalized_month, month_param, year_param, version_id, engine)
        loop_month = details['date'][5:7]
        if int(loop_month) > int(month_param):
            continue
        query_date = f"{year_param}-{details['date'][5:7]}-{details['date'][8:]}"
        replacements = {"{{date}}": query_date, "{version_id}": str(version)}
        sql_path = os.path.join(base_path, "sql_files/collapsible_items_query.sql")

        df = pd.read_sql_query(load_sql_query(sql_path, replacements), engine)
        df["real_value"] = df["effective_real_value"]
        populate_data_to_sheet(ws, df, type_id_to_rows, details["start_col"], details["sequence"])

        # --- Process Annual Data ---
        if annual_mapping_path and os.path.exists(annual_mapping_path):
            annual_mapping = read_json_mapping(annual_mapping_path).get("annual", {})
            if annual_mapping:
                start_col = annual_mapping["start_col"]
                sequence = annual_mapping["sequence"]
                query_date = f"{year_param}-01-01"
                replacements = {"2024-01-01": query_date, "{version_id}": str(version)}
                sql_path = os.path.join(base_path, "sql_files/collapsible_annual_query.sql")

                df_annual = pd.read_sql_query(load_sql_query(sql_path, replacements), engine)
                df["real_value"] = df["effective_real_value"]
                populate_data_to_sheet(ws, df_annual, type_id_to_rows, start_col, sequence)
    if cumul_mapping_path and os.path.exists(cumul_mapping_path):
        normalized_month = normalize_month_key(month_key)
        version = get_version_for_tbg_export(normalized_month, month_param, year_param, version_id, engine)

        cumul_mapping = read_json_mapping(cumul_mapping_path).get("cumul", {})
        if cumul_mapping:
            start_col = cumul_mapping["start_col"]
            value = ""
            if 4<=int(month_param)<=5:
                value = '1'
            elif 6<=int(month_param)<=8:
                value = '2'
            elif 9<=int(month_param)<=12:
                value = '3'
            sequence = cumul_mapping["sequence"]
            sequence = [ ele.replace('#', value) if '#' in ele else ele for ele in sequence]
            query_date = f"{year_param}-{month_param}-01"
            replacements = {'{{date}}': query_date, "{version_id}": str(version_id)}
            sql_path = os.path.join(base_path, "sql_files/collapsible_cumul_query.sql")

            df_cumul = pd.read_sql_query(load_sql_query(sql_path, replacements), engine)
            df["real_value"] = df["effective_real_value"]
            populate_data_to_sheet(ws, df_cumul, type_id_to_rows, start_col, sequence)


def main():
    args = parse_arguments()
    month_year = args.month_year
    version_id = args.version_id
    month_param, year_param = month_year[:2], month_year[2:]


    base_path = "benin_tbg_export"
    output_path = f"{base_path}/outputs/tbg_report_{month_year}.xlsx"
    config_path = os.path.join(base_path, "configs", "collapsible_config.json")

    sheets_to_process = read_json_mapping(config_path)
    wb = load_workbook(output_path)
    engine = get_db_engine()

    for sheet_data in sheets_to_process:
        process_sheet(sheet_data, wb, engine, year_param, month_param, version_id, base_path)

    wb.save(output_path)
    print("✅ Financial report generated for all months and annual data.")


if __name__ == "__main__":
    main()
