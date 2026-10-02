import os
import json
import pandas as pd
from openpyxl import load_workbook
from helpers.parse_arg import parse_arguments
from helpers.export_utils import excel_column_name, normalize_key, read_json_mapping,header_row_updater
from helpers.db_utils import get_db_engine
def actu_row_hider(sheet_name, ws):
    if sheet_name == "Data Mobile":
        ws.column_dimensions["CU"].hidden = True
        ws.column_dimensions["CW"].hidden = True
def excel_col_to_index(col_str):
    index = 0
    for i, c in enumerate(reversed(col_str.upper())):
        index += (ord(c) - 64) * (26 ** i)
    return index

def get_value_based_on_sheet(sheet_name, row_number, row, mapped_column):
    value = row.get(mapped_column)
    if value is None:
        return None

    ecart_cols = ['ecart_actual1_value', 'ecart_actual2_value', 'ecart_actual3_value', 'ecart_budget']

    if sheet_name == 'P&L conso' and row_number in (42, 56):
        if mapped_column in ecart_cols:
            return value * 100

    if sheet_name == 'Marge brute Mobile' and row_number == 61:
        if mapped_column in ecart_cols:
            return value * 100

    return value

def resolve_row_number(row_mapping, f_type, metric, submetric):
    for type_key, type_val in row_mapping.items():
        if normalize_key(type_key) != f_type:
            continue
        metrics = type_val.get("metrics", {})

        if metric != "n/a" and submetric != "n/a":
            for metric_key, metric_val in metrics.items():
                if normalize_key(metric_key) == metric:
                    submetrics = metric_val.get("submetrics", {})
                    for sub_key, row in submetrics.items():
                        if normalize_key(sub_key) == submetric:
                            return row

        if metric != "n/a":
            for metric_key, metric_val in metrics.items():
                if normalize_key(metric_key) == metric:
                    return metric_val.get("row_number")

        return type_val.get("row_number")
def mobile_and_data_header(ws, sheet_name, year_param):
    last_year = int(year_param) - 1

    if sheet_name == "Mobile Money":
        ws["CW4"] = last_year
    elif sheet_name == "Data Mobile":
        ws["CX4"] = last_year
    return
def update_cumul_headers_from_mapping_multiline(ws, cumul_config, year_param, sheet_name):


    current_year = str(int(year_param))
    last_year = str(int(year_param) - 1)

    sheets_with_last_year_only_evol = {"Data Mobile", "Mobile Money"}

    start_col = cumul_config["start_col"]
    header_row = cumul_config["header_row_index"]
    sequence = cumul_config["sequence"]

    start_col_index = excel_col_to_index(start_col)

    for i, col_type in enumerate(sequence):
        col_letter = excel_column_name(start_col_index + i)
        cell = ws[f"{col_letter}{header_row}"]

        if cell.value is None:
            continue

        text = str(cell.value)
        lines = text.split("\n")
        label = lines[1] if len(lines) > 1 else text  # keeps "Réel", "Budget", "ActuX", etc.

        if col_type == "real_value":
            if sheet_name in sheets_with_last_year_only_evol:
                cell.value = f"{current_year}"
            else:
                cell.value = f"{current_year}\nRéel"

        elif col_type == "budget_value":
            if sheet_name in sheets_with_last_year_only_evol:
                cell.value = f"{current_year}"
            else:
                cell.value = f"{current_year}\nBudget"

        elif col_type == "last_year_real_value":
            if sheet_name in sheets_with_last_year_only_evol:
                cell.value = f"{current_year}"
            else:
                cell.value = f"{last_year}\nRéel"

        elif col_type.startswith("actual"):
            cell.value = f"{current_year}\n{label}"
        # ecart_* intentionally untouched

def update_actual_column(ws, sheet_name,month_param, year_param,cell_address):
    month_param = int(month_param)
    if 4<=month_param<=5:
        if sheet_name == "Data Mobile":
            ws[cell_address[0]] = "Actu1"
            ws[cell_address[1]] = "Actu1"
            return
        ws[cell_address[0]] = f"{year_param }\nActu1"
        ws[cell_address[1]] = "Ecart /\nActu1"
    elif 6<=month_param<=8:
        if sheet_name == "Data Mobile":
            ws[cell_address[0]] = "Actu2"
            ws[cell_address[1]] = "Actu2"
            return
        ws[cell_address[0]] = f"{year_param }\nActu2"
        ws[cell_address[1]] = "Ecart /\nActu2"
    elif 9<=month_param<=12:
        if sheet_name == "Data Mobile":
            ws[cell_address[0]] = "Actu3"
            ws[cell_address[1]] = "Actu3"
            return
        ws[cell_address[0]] = f"{year_param }\nActu3"
        ws[cell_address[1]] = "Ecart /\nActu3"
def update_actual_column_name(ws, sheet_name, month_param, year_param):
    if sheet_name == "Mobile Money":
        return
    elif sheet_name == "Data Mobile":
        update_actual_column(ws, sheet_name,month_param,year_param, ["CU3","CW4"])
    else:
        update_actual_column(ws,sheet_name,month_param,year_param, ["CT5","CV5"])
def process_cumulative_sheet(sheet, wb, engine, year_param, month_param, version_id, base_path, is_calculated_by_formula):
    date_param = f"{year_param}-{month_param}-01"
    sheet_name = sheet["sheet_name"].strip()
    sheet_name = {
        "Marge brute Mobile": "Marge Mobile",
        "Parc Mobile": "Parc Mobile ",
    }.get(sheet_name, sheet_name)
    ws = wb[sheet_name]

    with open(sheet["sheet_mapping_path"],) as f:
        row_mapping = json.load(f)

    with open(sheet["cumul_mapping_path"],) as f:
        cumul_config = json.load(f)["cumul"]

    start_col_index = excel_col_to_index(cumul_config["start_col"])
    sequence = cumul_config["sequence"]
    header_row_index = cumul_config["header_row_index"]
    header_row_updater(ws,sequence,cumul_config["start_col"], year_param, header_row_index)
    update_cumul_headers_from_mapping_multiline(ws, cumul_config, year_param, sheet_name)
    mobile_and_data_header(ws,sheet_name, year_param )
    update_actual_column_name(ws, sheet_name, month_param, year_param)
    if is_calculated_by_formula:
        return
    with open(f"{base_path}/sql_files/cumul_trafic_mobile_query.sql") as f:
        query = f.read()

    sheet_name = sheet["sheet_name"]
    if sheet_name == "Cash conso":
        sheet_name = "Cash Conso"
    df = pd.read_sql_query(query, engine, params={
            "sheet_name": sheet_name,
            "date_param": date_param,
            "version_id": version_id
        })
    for _, row in df.iterrows():
        f_type = normalize_key(row["financial_type_name"])
        metric = normalize_key(row["metric_name"])
        submetric = normalize_key(row["submetric_name"])
        row_number = resolve_row_number(row_mapping, f_type, metric, submetric)

        if not row_number:
            print(f"⚠️ Skipped: {row['financial_type_name']} → {row['metric_name']} → {row['submetric_name']}")
            continue
        value = ""
        if 4<=int(month_param)<=5:
            value = '1'
        elif 6<=int(month_param)<=8:
            value = '2'
        elif 9<=int(month_param)<=12:
            value = '3'
        sequence = [ ele.replace('#', value) if '#' in ele else ele for ele in sequence]
        for i, metric_key in enumerate(sequence):
            if metric_key in row:
                col_letter = excel_column_name(start_col_index + i)
                cell_address = f"{col_letter}{row_number}"
                ws[cell_address] = get_value_based_on_sheet(sheet_name, row_number, row, metric_key)

    actu_row_hider(sheet_name, ws)

def main():
    args = parse_arguments()
    month_year = args.month_year
    version_id = args.version_id
    month_param, year_param = month_year[:2], month_year[2:]

    base_path = "benin_tbg_export"
    output_path = f"{base_path}/outputs/tbg_report_{month_year}.xlsx"

    wb = load_workbook(output_path)
    sheets_to_process = read_json_mapping(f"{base_path}/configs/config.json")
    engine = get_db_engine()

    try:
        for sheet in sheets_to_process:
            print(f"Processing sheet '{sheet['sheet_name']}'")
            process_cumulative_sheet(sheet, wb, engine, year_param, month_param, version_id, base_path, sheet["is_calculated_by_formula"])

        wb.save(output_path)
        print(f"\n✅ Financial report generated: {output_path}")

    except Exception as e:
        print(f"[ERROR] {str(e)}")

if __name__ == "__main__":
    main()
